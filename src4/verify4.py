# -*- coding: utf-8 -*-
"""Kiểm tra cấu trúc, tính nhất quán và chất lượng hình thức của 20 tệp .docx.

Chạy sau build4.py. Báo LỖI cho những vấn đề bắt buộc phải sửa và CẢNH BÁO cho
những điểm cần rà lại bằng mắt.
"""
import os
import re
import sys
import zipfile
import difflib
from collections import Counter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
FIGS = os.path.join(ROOT, "figs")
OUTDIR = os.path.join(ROOT, "DE_KIEM_TRA_THEO_CHUONG")
sys.path.insert(0, HERE)

from docx import Document
import de_ch1
import de_ch2
import build4

ERR, WARN = [], []
LETTERS = ["A", "B", "C", "D"]
ALL_TESTS = [(1, t) for t in de_ch1.DE_CH1] + [(2, t) for t in de_ch2.DE_CH2]

# Áp dụng đúng phép cân bằng vị trí đáp án mà build4 đã dùng, để kiểm tra trên
# chính dữ liệu đã được xuất ra tệp .docx.
for _ch, _tests in ((1, de_ch1.DE_CH1), (2, de_ch2.DE_CH2)):
    for _k, _t in enumerate(_tests):
        build4.balance_test(_t, seed=1000 * _ch + _k, offset=(_ch - 1) * 5 + _k)
        build4.thin_figures(_t)

LEAK = ["Đáp án:", "Lời giải:", "Giải thích:", "→ ĐÚNG", "→ SAI", "BẢNG ĐÁP ÁN"]
LETTER_REF = re.compile(r"(?:phương án|đáp án|Đáp án|Phương án)\s+[ABCD]\b")
SHORT_OK = re.compile(r"^[−-]?\d+(?:,\d+)?$")


def err(m):
    ERR.append(m)


def warn(m):
    WARN.append(m)


def docx_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.extend(p.text for p in c.paragraphs)
    return "\n".join(parts)


def n_images(path):
    with zipfile.ZipFile(path) as z:
        return sum(1 for n in z.namelist() if n.startswith("word/media/"))


# ------------------------------------------------------ 1. cấu trúc dữ liệu
print("1) CẤU TRÚC ĐỀ")
used_figs = set()
for ch, t in ALL_TESTS:
    code = t["code"]
    if len(t["p1"]) != 18:
        err("%s: Phần I có %d câu (phải là 18)" % (code, len(t["p1"])))
    if len(t["p2"]) != 4:
        err("%s: Phần II có %d câu (phải là 4)" % (code, len(t["p2"])))
    if len(t["p3"]) != 6:
        err("%s: Phần III có %d câu (phải là 6)" % (code, len(t["p3"])))

    for i, it in enumerate(t["p1"], 1):
        if len(it["o"]) != 4:
            err("%s Phần I câu %d: có %d phương án" % (code, i, len(it["o"])))
        if len(set(o.strip().lower() for o in it["o"])) != 4:
            err("%s Phần I câu %d: có phương án trùng nhau" % (code, i))
        if any(not o.strip() for o in it["o"]):
            err("%s Phần I câu %d: có phương án rỗng" % (code, i))
        if it["a"] not in LETTERS:
            err("%s Phần I câu %d: khoá đáp án không hợp lệ" % (code, i))
        if len(it.get("sol", "")) < 60:
            warn("%s Phần I câu %d: lời giải quá ngắn (%d kí tự)"
                 % (code, i, len(it.get("sol", ""))))
        if it.get("fig") or it.get("_fig_hidden"):
            used_figs.add(it.get("fig") or it["_fig_hidden"])

    for i, it in enumerate(t["p2"], 1):
        if len(it["items"]) != 4:
            err("%s Phần II câu %d: có %d ý (phải là 4)" % (code, i, len(it["items"])))
        vals = [v for _x, v, _e in it["items"]]
        if all(vals) or not any(vals):
            warn("%s Phần II câu %d: cả bốn ý cùng đúng hoặc cùng sai" % (code, i))
        for lab, (_x, _v, ex) in zip("abcd", it["items"]):
            if len(ex) < 60:
                warn("%s Phần II câu %d ý %s: giải thích quá ngắn" % (code, i, lab))
        if it.get("fig") or it.get("_fig_hidden"):
            used_figs.add(it.get("fig") or it["_fig_hidden"])

    for i, it in enumerate(t["p3"], 1):
        a = it["ans"].strip()
        if not SHORT_OK.match(a):
            err("%s Phần III câu %d: đáp án “%s” không phải dạng số điền được"
                % (code, i, a))
        if len(a) > 5:
            err("%s Phần III câu %d: đáp án “%s” dài quá 4–5 kí tự ô tô"
                % (code, i, a))
        if len(it.get("sol", "")) < 60:
            warn("%s Phần III câu %d: lời giải quá ngắn" % (code, i))
        if it.get("fig") or it.get("_fig_hidden"):
            used_figs.add(it.get("fig") or it["_fig_hidden"])
print("   đã kiểm tra %d đề × 28 câu = %d câu"
      % (len(ALL_TESTS), 28 * len(ALL_TESTS)))

# ------------------------------------------- 2. lời giải không dẫn chiếu chữ cái
print("2) LỜI GIẢI KHÔNG DẪN CHIẾU CHỮ CÁI PHƯƠNG ÁN")
nref = 0
for ch, t in ALL_TESTS:
    texts = [(("P1 c%d" % i), it["sol"]) for i, it in enumerate(t["p1"], 1)]
    texts += [(("P3 c%d" % i), it["sol"]) for i, it in enumerate(t["p3"], 1)]
    for i, it in enumerate(t["p2"], 1):
        for lab, (_x, _v, ex) in zip("abcd", it["items"]):
            texts.append(("P2 c%d%s" % (i, lab), ex))
    for where, s in texts:
        for m in LETTER_REF.finditer(s):
            err("%s %s: lời giải dẫn chiếu chữ cái “%s”"
                % (t["code"], where, m.group(0)))
            nref += 1
print("   số chỗ dẫn chiếu chữ cái: %d" % nref)

# --------------------------------------------- 3. trùng lặp câu hỏi giữa các đề
print("3) TRÙNG LẶP PHẦN DẪN GIỮA CÁC CÂU")
stems = []
for ch, t in ALL_TESTS:
    for i, it in enumerate(t["p1"], 1):
        stems.append(("%s P1.%d" % (t["code"], i), it["q"]))
    for i, it in enumerate(t["p2"], 1):
        stems.append(("%s P2.%d" % (t["code"], i), it["stem"]))
    for i, it in enumerate(t["p3"], 1):
        stems.append(("%s P3.%d" % (t["code"], i), it["q"]))
ndup = 0
for a in range(len(stems)):
    for b in range(a + 1, len(stems)):
        r = difflib.SequenceMatcher(None, stems[a][1], stems[b][1]).ratio()
        if r > 0.88:
            warn("trùng %.2f: %s  ↔  %s" % (r, stems[a][0], stems[b][0]))
            ndup += 1
print("   đã so sánh %d cặp, phát hiện %d cặp giống trên 0,88"
      % (len(stems) * (len(stems) - 1) // 2, ndup))

# --------------------------------------------------- 4. thiên lệch độ dài đáp án
print("4) ĐỘ DÀI PHƯƠNG ÁN ĐÚNG SO VỚI PHƯƠNG ÁN NHIỄU")
longest, total, gaps, gross = 0, 0, [], 0
for ch, t in ALL_TESTS:
    for i, it in enumerate(t["p1"], 1):
        key = it["o"][LETTERS.index(it["a"])]
        others = [o for o in it["o"] if o is not key]
        mx = max(len(o) for o in others)
        total += 1
        if len(key) > mx:
            longest += 1
        gaps.append(len(key) - sum(len(o) for o in others) / 3.0)
        if len(key) > mx + 20 and len(key) > 1.35 * mx:
            gross += 1
            warn("%s P1.%d: phương án đúng dài hơn hẳn các phương án nhiễu"
                 % (t["code"], i))
print("   phương án đúng là dài nhất: %d/%d câu (%.0f%%; ngẫu nhiên ≈ 25%%)"
      % (longest, total, 100.0 * longest / total))
print("   chênh lệch độ dài trung bình: %+.1f kí tự  |  câu lệch mạnh: %d"
      % (sum(gaps) / len(gaps), gross))
if longest / total > 0.42:
    warn("tỉ lệ “đáp án đúng dài nhất” cao bất thường, học sinh có thể đoán mò")

# ------------------------------------------------------------- 5. phân bố đáp án
print("5) PHÂN BỐ VỊ TRÍ ĐÁP ÁN PHẦN I")
allc = Counter()
for ch, t in ALL_TESTS:
    c = Counter(it["a"] for it in t["p1"])
    allc.update(c)
    if max(c.values()) > 7:
        warn("%s: có %d câu cùng đáp án %s"
             % (t["code"], max(c.values()), c.most_common(1)[0][0]))
print("   toàn bộ 10 đề:", dict(sorted(allc.items())))

# ------------------------------------------------------------------- 6. hình vẽ
print("6) HÌNH VẼ")
for name in sorted(used_figs):
    if not os.path.exists(os.path.join(FIGS, name + ".png")):
        err("thiếu tệp hình %s.png" % name)
drawn = {f[:-4] for f in os.listdir(FIGS) if f.startswith("t") and f.endswith(".png")
         and re.match(r"^t[12][1-5][a-d]\.png$", f)}
print("   số hình đã vẽ cho bộ đề: %d  |  số hình được dùng: %d"
      % (len(drawn), len(used_figs)))
for name in sorted(drawn - used_figs):
    warn("hình %s đã vẽ nhưng không được dùng" % name)

# ---------------------------------------------------- 7. kiểm tra các tệp .docx
print("7) TỆP .DOCX")
if not os.path.isdir(OUTDIR):
    err("chưa có thư mục %s — hãy chạy build4.py trước" % OUTDIR)
else:
    files = sorted(os.listdir(OUTDIR))
    exams = [f for f in files if f.endswith(".docx") and "_Loi_giai" not in f]
    sols = [f for f in files if f.endswith("_Loi_giai.docx")]
    if len(exams) != 10 or len(sols) != 10:
        err("có %d đề và %d lời giải (phải là 10 và 10)" % (len(exams), len(sols)))
    for e in exams:
        if e.replace(".docx", "_Loi_giai.docx") not in sols:
            err("đề %s không có tệp lời giải tương ứng" % e)

    for ch, t in ALL_TESTS:
        base = "Chuong_%d_De_%s" % (ch, t["so"])
        pe = os.path.join(OUTDIR, base + ".docx")
        ps = os.path.join(OUTDIR, base + "_Loi_giai.docx")
        if not (os.path.exists(pe) and os.path.exists(ps)):
            err("thiếu tệp cho %s" % base)
            continue
        te, ts = docx_text(pe), docx_text(ps)

        # 7a. đề không được lộ đáp án
        for k in LEAK:
            if k in te:
                err("%s: tệp đề chứa dấu hiệu lộ đáp án “%s”" % (base, k))
        # 7b. đề không chứa nhãn mức độ / chủ đề nội bộ
        for k in ["Nhận biết", "Thông hiểu", "Vận dụng cao", "Mức 1", "Mức 4"]:
            if k in te:
                err("%s: tệp đề chứa nhãn nội bộ “%s”" % (base, k))
        # 7c. đồng bộ nội dung câu hỏi giữa đề và lời giải
        miss = 0
        for it in t["p1"]:
            probe = it["q"][:55]
            if probe not in te:
                err("%s: đề thiếu câu “%s…”" % (base, probe[:35])); miss += 1
            if probe not in ts:
                err("%s: lời giải thiếu câu “%s…”" % (base, probe[:35])); miss += 1
        for it in t["p2"]:
            probe = it["stem"][:55]
            if probe not in te or probe not in ts:
                err("%s: lệch phần II “%s…”" % (base, probe[:35])); miss += 1
        for it in t["p3"]:
            probe = it["q"][:55]
            if probe not in te or probe not in ts:
                err("%s: lệch phần III “%s…”" % (base, probe[:35])); miss += 1
        # 7d. lời giải phải có đủ đáp án
        for it in t["p3"]:
            if it["ans"] not in ts:
                err("%s: lời giải thiếu đáp án “%s”" % (base, it["ans"]))
        # 7e. số hình nhúng
        nfig = len({it.get("fig") for lst in (t["p1"], t["p2"], t["p3"])
                    for it in lst if it.get("fig")})
        ne, ns = n_images(pe), n_images(ps)
        print("   %-22s  %2d hình khác nhau | đề %2d ảnh, lời giải %2d ảnh | "
              "%4d KB + %4d KB"
              % (base, nfig, ne, ns,
                 os.path.getsize(pe) // 1024, os.path.getsize(ps) // 1024))
        if ne == 0:
            warn("%s: tệp đề không có hình nào" % base)

# --------------------------------------------------- 8. vệ sinh văn bản
print("8) VỆ SINH VĂN BẢN")
nbad = 0
for ch, t in ALL_TESTS:
    texts = []
    for i, it in enumerate(t["p1"], 1):
        texts.append(("P1.%d dẫn" % i, it["q"]))
        for j, o in enumerate(it["o"]):
            texts.append(("P1.%d %s" % (i, LETTERS[j]), o))
    for i, it in enumerate(t["p2"], 1):
        texts.append(("P2.%d dẫn" % i, it["stem"]))
        for lab, (x, _v, _e) in zip("abcd", it["items"]):
            texts.append(("P2.%d%s" % (i, lab), x))
    for i, it in enumerate(t["p3"], 1):
        texts.append(("P3.%d dẫn" % i, it["q"]))
    for where, txt in texts:
        if "  " in txt or txt != txt.strip():
            err("%s %s: có khoảng trắng thừa" % (t["code"], where)); nbad += 1
        if "\n" in txt:
            err("%s %s: phần đề bài chứa dấu xuống dòng bất thường"
                % (t["code"], where)); nbad += 1
        if txt.count("“") != txt.count("”"):
            err("%s %s: dấu ngoặc kép không cân" % (t["code"], where)); nbad += 1
        if "'" in txt:
            warn("%s %s: dùng dấu nháy thẳng thay cho ngoặc kép cong"
                 % (t["code"], where))
print("   đã soi %d chuỗi đề bài, số chỗ bất thường: %d"
      % (sum(1 for _ in range(0)) or 10 * (18 * 5 + 4 * 5 + 6), nbad))

# ------------------------------------------- 9. tiến trình độ khó giữa các đề
print("9) TIẾN TRÌNH ĐỘ KHÓ (đo bằng tỉ lệ câu định lượng)")
NUM = re.compile(r"\d")
for ch, t in ALL_TESTS:
    ncalc = sum(1 for it in t["p1"] if NUM.search(it["q"]) or
                any(NUM.search(o) for o in it["o"]))
    nds = sum(1 for it in t["p2"]
              if any(NUM.search(x) for x, _v, _e in it["items"]))
    nlong = sum(len(it["sol"]) for it in t["p1"]) // len(t["p1"])
    print("   %s: %2d/18 câu Phần I có tính toán | %d/4 câu Phần II có số liệu | "
          "lời giải trung bình %d kí tự" % (t["code"], ncalc, nds, nlong))

# ------------------------------------------------------------------ tổng kết
print()
print("=" * 74)
for w in WARN:
    print("  ! CẢNH BÁO:", w)
print("  CẢNH BÁO: %d" % len(WARN))
for e in ERR:
    print("  ✗ LỖI:", e)
print("  LỖI: %d" % len(ERR))
if not ERR:
    print("  → KHÔNG CÓ LỖI CẤU TRÚC.")
