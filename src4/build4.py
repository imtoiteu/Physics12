# -*- coding: utf-8 -*-
"""Sinh 20 tệp Word: 10 đề kiểm tra theo chương + 10 tệp lời giải chi tiết.

Mỗi đề mô phỏng đúng cấu trúc đề thi tốt nghiệp THPT môn Vật lí hiện hành
(Quyết định 764/QĐ-BGDĐT ngày 08/3/2024, áp dụng từ năm 2025 và giữ ổn định
cho các năm sau): 28 câu / 40 lệnh hỏi / 50 phút / thang điểm 10.
"""

import os
import re
import sys
import random

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
FIGS = os.path.join(ROOT, "figs")
OUTDIR = os.path.join(ROOT, "DE_KIEM_TRA_THEO_CHUONG")
sys.path.insert(0, HERE)

import de_ch1
import de_ch2

LETTERS = ["A", "B", "C", "D"]
DS_LABELS = ["a)", "b)", "c)", "d)"]
FONT = "Times New Roman"
DARKRED = RGBColor(0xA6, 0x1B, 0x1B)
DARKBLUE = RGBColor(0x1F, 0x3A, 0x6E)
DARKGREEN = RGBColor(0x1E, 0x64, 0x36)
GREY = RGBColor(0x55, 0x55, 0x55)

# Kí hiệu dạng "k_B", "Q_toả", "c_nước"… được in thành chỉ số dưới thật sự trong Word.
SUBSCRIPT = re.compile(r"([A-Za-zΔρλ])_((?:[^\W_]){1,8})", re.UNICODE)

THOI_GIAN = 50
PART_TITLES = [
    ("PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.",
     "Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án."),
    ("PHẦN II. Câu trắc nghiệm đúng sai.",
     "Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh "
     "chọn đúng hoặc sai."),
    ("PHẦN III. Câu trắc nghiệm trả lời ngắn.",
     "Thí sinh trả lời từ câu 1 đến câu 6."),
]

CACH_TINH_DIEM = (
    "• Phần I: mỗi câu trả lời đúng được 0,25 điểm (18 câu → 4,5 điểm).\n"
    "• Phần II: mỗi câu gồm 4 ý; đúng 1 ý được 0,10 điểm, đúng 2 ý được 0,25 điểm, đúng "
    "3 ý được 0,50 điểm, đúng cả 4 ý được 1,00 điểm (4 câu → 4,0 điểm).\n"
    "• Phần III: mỗi câu trả lời đúng được 0,25 điểm (6 câu → 1,5 điểm).\n"
    "• Tổng điểm toàn bài: 10,0 điểm."
)


# --------------------------------------------------------------------- tiện ích
def set_base(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.line_spacing = 1.05
    for s in doc.sections:
        s.top_margin = Cm(1.5); s.bottom_margin = Cm(1.4)
        s.left_margin = Cm(1.8); s.right_margin = Cm(1.4)


def emit(p, text, bold=False, italic=False, size=12, color=None):
    """Thêm text vào đoạn p, tự động chuyển kí hiệu dạng X_abc thành chỉ số dưới."""
    pos = 0
    for m in SUBSCRIPT.finditer(text):
        for chunk, sub in ((text[pos:m.start()], False), (m.group(1), False),
                           (m.group(2), True)):
            if not chunk:
                continue
            r = p.add_run(chunk)
            r.bold = bold; r.italic = italic
            r.font.size = Pt(size * (0.78 if sub else 1.0)); r.font.name = FONT
            r.font.subscript = sub
            if color is not None:
                r.font.color.rgb = color
        pos = m.end()
    tail = text[pos:]
    if tail or pos == 0:
        r = p.add_run(tail)
        r.bold = bold; r.italic = italic
        r.font.size = Pt(size); r.font.name = FONT
        if color is not None:
            r.font.color.rgb = color


def P(doc, text="", bold=False, italic=False, size=12, align=None, before=0, after=3,
      indent=None, color=None, first_line=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if indent is not None:
        pf.left_indent = Cm(indent)
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)
    if text:
        for i, line in enumerate(str(text).split("\n")):
            if i:
                p.add_run().add_break()
            emit(p, line, bold, italic, size, color)
    return p


def rich(doc, parts, indent=None, before=0, after=3, size=12, align=None):
    """parts = [(text, bold, italic, color|None), ...]"""
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if indent is not None:
        pf.left_indent = Cm(indent)
    for text, b, i, c in parts:
        for k, line in enumerate(str(text).split("\n")):
            if k:
                p.add_run().add_break()
            emit(p, line, b, i, size, c)
    return p


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def no_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        borders.append(el)
    tblPr.append(borders)


def box(doc, title, text, fill="EEF3FA", tcolor=DARKBLUE, size=11.5):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    shade(c, fill)
    c.paragraphs[0].paragraph_format.space_after = Pt(2)
    if title:
        r = c.paragraphs[0].add_run(title)
        r.bold = True; r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = tcolor
        tp = c.add_paragraph()
    else:
        tp = c.paragraphs[0]
    tp.paragraph_format.space_after = Pt(1)
    for i, line in enumerate(str(text).split("\n")):
        if i:
            tp.add_run().add_break()
        emit(tp, line, size=size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def picture(doc, name, caption=None, maxw=9.6, maxh=5.4):
    path = os.path.join(FIGS, name + ".png")
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    with Image.open(path) as im:
        w, h = im.size
    width = maxw
    if h / w * width > maxh:
        width = maxh * w / h
    doc.add_picture(path, width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.space_before = Pt(3)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(1)
    if caption:
        P(doc, caption, italic=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=5)


def simple_table(doc, headers, rows, size=10.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, htxt in enumerate(headers):
        c = t.rows[0].cells[j]
        shade(c, "DCE6F1")
        c.paragraphs[0].paragraph_format.space_after = Pt(1)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.paragraphs[0].add_run(htxt)
        r.bold = True; r.font.size = Pt(size); r.font.name = FONT
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].paragraphs[0].paragraph_format.space_after = Pt(1)
            cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cells[j].paragraphs[0].add_run(str(val))
            r.font.size = Pt(size); r.font.name = FONT
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return t


# ------------------------------------------------------- cân bằng vị trí đáp án
def balance_test(test, seed, offset=0):
    """Hoán vị vị trí các phương án để đáp án đúng phân bố đều giữa A, B, C, D.

    Nội dung các phương án hoàn toàn không đổi, chỉ đổi chỗ cho nhau. Vì 18 không
    chia hết cho 4, mỗi đề dư hai chữ cái; tham số offset làm hai chữ cái dư đó
    xoay vòng giữa các đề để tổng thể vẫn cân bằng."""
    rng = random.Random(seed)
    n = len(test["p1"])
    targets = [LETTERS[(i + offset) % 4] for i in range(n)]
    for _ in range(500):
        rng.shuffle(targets)
        if not any(targets[i] == targets[i + 1] == targets[i + 2]
                   for i in range(n - 2)):
            break
    for item, dst in zip(test["p1"], targets):
        cur = item["a"]
        if cur == dst:
            continue
        i, j = LETTERS.index(cur), LETTERS.index(dst)
        item["o"][i], item["o"][j] = item["o"][j], item["o"][i]
        item["a"] = dst


# --------------------------------------------------------------- hiển thị câu
def _extras(doc, item, maxw=9.6, maxh=5.4):
    if item.get("fig"):
        picture(doc, item["fig"], maxw=maxw, maxh=maxh)
    if item.get("tbl"):
        cap, hdr, rows = item["tbl"]
        if cap:
            P(doc, cap, italic=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER,
              before=3, after=2)
        simple_table(doc, hdr, rows)


def options_layout(doc, opts):
    """Bố trí 4 phương án: 4 cột / 2 cột / mỗi phương án một dòng, tuỳ độ dài."""
    longest = max(len(o) for o in opts)
    if longest <= 16:
        ncol = 4
    elif longest <= 38:
        ncol = 2
    else:
        ncol = 1
    if ncol == 1:
        for i, opt in enumerate(opts):
            rich(doc, [("%s. " % LETTERS[i], True, False, None),
                       (opt, False, False, None)], indent=0.6, after=0)
        return
    nrow = 4 // ncol
    t = doc.add_table(rows=nrow, cols=ncol)
    no_borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, opt in enumerate(opts):
        cell = t.rows[i // ncol].cells[i % ncol]
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run("%s. " % LETTERS[i]); r.bold = True
        r.font.size = Pt(12); r.font.name = FONT
        emit(p, opt, size=12)


def keep(*paras):
    """Không cho Word ngắt trang ở giữa một câu hỏi."""
    for p in paras:
        if p is None:
            continue
        p.paragraph_format.keep_together = True
        p.paragraph_format.keep_with_next = True


def render_mc(doc, no, item):
    first = len(doc.paragraphs)
    rich(doc, [("Câu %d. " % no, True, False, None), (item["q"], False, False, None)],
         before=5, after=1)
    _extras(doc, item)
    options_layout(doc, item["o"])
    keep(*doc.paragraphs[first:-1])


def render_ds(doc, no, item, show=False):
    first = len(doc.paragraphs)
    rich(doc, [("Câu %d. " % no, True, False, None), (item["stem"], False, False, None)],
         before=7, after=1)
    _extras(doc, item)
    for lab, (txt, val, _ex) in zip(DS_LABELS, item["items"]):
        parts = [(lab + " ", True, False, None), (txt, False, False, None)]
        if show:
            parts.append(("   → " + ("ĐÚNG" if val else "SAI"), True, False,
                          DARKBLUE if val else DARKRED))
        rich(doc, parts, indent=0.6, after=1)
    keep(*doc.paragraphs[first:-1])


def render_short(doc, no, item):
    first = len(doc.paragraphs)
    rich(doc, [("Câu %d. " % no, True, False, None), (item["q"], False, False, None)],
         before=7, after=1)
    _extras(doc, item)
    P(doc, "Trả lời: ................................................",
      size=12, indent=0.6, after=1)
    keep(*doc.paragraphs[first:-1])


# ------------------------------------------------------------------ đầu đề thi
def exam_header(doc, test, solution=False):
    t = doc.add_table(rows=1, cols=2)
    no_borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = t.rows[0].cells

    def put(cell, lines):
        first = True
        for txt, bold, size in lines:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(txt)
            r.bold = bold; r.font.size = Pt(size); r.font.name = FONT

    put(left, [("BỘ ĐỀ KIỂM TRA THEO CHƯƠNG", True, 11.5),
               ("VẬT LÍ 12 – CHƯƠNG TRÌNH GDPT 2018", True, 11.5),
               ("———————", False, 11)])
    put(right, [("ĐỀ KIỂM TRA ĐÁNH GIÁ NĂNG LỰC", True, 11.5),
                ("Môn: VẬT LÍ – Lớp 12", True, 11.5),
                ("Thời gian làm bài: %d phút" % THOI_GIAN, False, 11)])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    P(doc, test["title"] + ("  –  HƯỚNG DẪN GIẢI CHI TIẾT" if solution else ""),
      bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER, after=1, color=DARKRED)
    P(doc, test["subtitle"], italic=True, size=11.5,
      align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    P(doc, "Mã đề: %s   |   Cấu trúc: 18 câu trắc nghiệm + 4 câu đúng/sai + 6 câu trả lời "
           "ngắn (28 câu, 40 lệnh hỏi)" % test["code"],
      size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=6, color=GREY)

    if not solution:
        t2 = doc.add_table(rows=1, cols=2)
        no_borders(t2)
        c1, c2 = t2.rows[0].cells
        p = c1.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
        r = p.add_run("Họ và tên học sinh: ................................."
                      "................")
        r.font.size = Pt(12); r.font.name = FONT
        p = c2.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("Ngày làm bài: ...../...../..........   Điểm: ..........")
        r.font.size = Pt(12); r.font.name = FONT
        doc.add_paragraph().paragraph_format.space_after = Pt(2)


def part_heading(doc, idx):
    title, guide = PART_TITLES[idx]
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.bold = True; r.font.size = Pt(13); r.font.name = FONT; r.font.color.rgb = DARKRED
    gp = P(doc, guide, italic=True, size=11, after=2)
    gp.paragraph_format.keep_with_next = True


# ------------------------------------------------------------------- ĐỀ THI
def thin_figures(test):
    """Bỏ hình lặp lại nếu cùng một hình vừa xuất hiện ở một trong hai câu liền
    trước trong cùng một phần: khi đó phần dẫn đã dẫn chiếu tới câu có hình."""
    for key in ("p1", "p2", "p3"):
        recent = []
        for it in test[key]:
            f = it.get("fig")
            if f and f in recent[-2:]:
                it["_fig_hidden"] = f
                it.pop("fig")
            recent.append(f)


def build_exam(test, hangso, path):
    doc = Document(); set_base(doc)
    exam_header(doc, test)
    box(doc, None, hangso, fill="F4F6F8", size=11)

    part_heading(doc, 0)
    for i, item in enumerate(test["p1"], 1):
        render_mc(doc, i, item)

    part_heading(doc, 1)
    for i, item in enumerate(test["p2"], 1):
        render_ds(doc, i, item)

    part_heading(doc, 2)
    for i, item in enumerate(test["p3"], 1):
        render_short(doc, i, item)

    P(doc, "----------- HẾT -----------", bold=True,
      align=WD_ALIGN_PARAGRAPH.CENTER, before=14, after=2)
    P(doc, "Học sinh không được sử dụng tài liệu. Cán bộ coi thi không giải thích gì thêm.",
      italic=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.save(path)
    return path


# ---------------------------------------------------------------- LỜI GIẢI
def ds_string(item):
    return "".join("Đ" if v else "S" for _t, v, _e in item["items"])


def answer_tables(doc, test):
    P(doc, "BẢNG ĐÁP ÁN", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER,
      before=6, after=4, color=DARKRED)

    P(doc, "Phần I – Trắc nghiệm nhiều phương án lựa chọn", bold=True, size=11.5,
      before=4, after=2, color=DARKBLUE)
    ids = [str(i) for i in range(1, 10)]
    ans = [test["p1"][i - 1]["a"] for i in range(1, 10)]
    simple_table(doc, ["Câu"] + ids, [["Đáp án"] + ans], size=10.5)
    ids = [str(i) for i in range(10, 19)]
    ans = [test["p1"][i - 1]["a"] for i in range(10, 19)]
    simple_table(doc, ["Câu"] + ids, [["Đáp án"] + ans], size=10.5)

    P(doc, "Phần II – Trắc nghiệm đúng/sai (ghi theo thứ tự bốn ý a, b, c, d)",
      bold=True, size=11.5, before=6, after=2, color=DARKBLUE)
    simple_table(doc, ["Câu", "1", "2", "3", "4"],
                 [["Đáp án"] + [ds_string(it) for it in test["p2"]]], size=10.5)

    P(doc, "Phần III – Trả lời ngắn", bold=True, size=11.5,
      before=6, after=2, color=DARKBLUE)
    simple_table(doc, ["Câu", "1", "2", "3", "4", "5", "6"],
                 [["Đáp án"] + [it["ans"] for it in test["p3"]]], size=10.5)


def build_solution(test, hangso, path):
    doc = Document(); set_base(doc)
    exam_header(doc, test, solution=True)
    box(doc, "CÁCH TÍNH ĐIỂM THEO QUY ĐỊNH HIỆN HÀNH", CACH_TINH_DIEM,
        fill="F2F8F0", tcolor=DARKGREEN, size=11)
    box(doc, None, hangso, fill="F4F6F8", size=11)
    answer_tables(doc, test)
    doc.add_page_break()

    P(doc, "PHẦN I – HƯỚNG DẪN GIẢI CHI TIẾT", bold=True, size=13.5,
      align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=5, color=DARKRED)
    for i, item in enumerate(test["p1"], 1):
        rich(doc, [("Câu %d. " % i, True, False, None), (item["q"], False, True, None)],
             before=8, after=2)
        _extras(doc, item, maxw=8.8, maxh=5.0)
        key = item["a"]
        rich(doc, [("Đáp án: %s. " % key, True, False, DARKRED),
                   (item["o"][LETTERS.index(key)], True, False, DARKRED)],
             indent=0.6, after=2)
        rich(doc, [("Lời giải: ", True, True, None),
                   (item["sol"], False, False, None)], indent=0.6, after=2)

    doc.add_page_break()
    P(doc, "PHẦN II – HƯỚNG DẪN GIẢI CHI TIẾT", bold=True, size=13.5,
      align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=5, color=DARKRED)
    for i, item in enumerate(test["p2"], 1):
        rich(doc, [("Câu %d. " % i, True, False, None),
                   (item["stem"], False, True, None)], before=9, after=2)
        _extras(doc, item, maxw=8.8, maxh=5.0)
        rich(doc, [("Đáp án: ", True, False, DARKRED),
                   (ds_string(item), True, False, DARKRED)], indent=0.6, after=2)
        for lab, (txt, val, ex) in zip(DS_LABELS, item["items"]):
            rich(doc, [(lab + " ", True, False, None), (txt, False, False, None),
                       ("  →  " + ("ĐÚNG" if val else "SAI"), True, False,
                        DARKBLUE if val else DARKRED)], indent=0.6, after=1)
            rich(doc, [("Giải thích: ", True, True, None), (ex, False, False, None)],
                 indent=1.1, after=3, size=11.5)

    doc.add_page_break()
    P(doc, "PHẦN III – HƯỚNG DẪN GIẢI CHI TIẾT", bold=True, size=13.5,
      align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=5, color=DARKRED)
    for i, item in enumerate(test["p3"], 1):
        rich(doc, [("Câu %d. " % i, True, False, None), (item["q"], False, True, None)],
             before=9, after=2)
        _extras(doc, item, maxw=8.8, maxh=5.0)
        rich(doc, [("Đáp án: ", True, False, DARKRED),
                   (item["ans"], True, False, DARKRED)], indent=0.6, after=2)
        rich(doc, [("Lời giải: ", True, True, None), (item["sol"], False, False, None)],
             indent=0.6, after=2)

    P(doc, "----------- HẾT -----------", bold=True,
      align=WD_ALIGN_PARAGRAPH.CENTER, before=14)
    doc.save(path)
    return path


# --------------------------------------------------------------------- chạy
def main():
    os.makedirs(OUTDIR, exist_ok=True)
    made = []
    counts = {L: 0 for L in LETTERS}
    for chuong, (tests, hangso) in enumerate(
            [(de_ch1.DE_CH1, de_ch1.HANG_SO), (de_ch2.DE_CH2, de_ch2.HANG_SO)], 1):
        for k, test in enumerate(tests):
            balance_test(test, seed=1000 * chuong + k, offset=(chuong - 1) * 5 + k)
            thin_figures(test)
            for it in test["p1"]:
                counts[it["a"]] += 1
            base = "Chuong_%d_De_%s" % (chuong, test["so"])
            p1 = build_exam(test, hangso, os.path.join(OUTDIR, base + ".docx"))
            p2 = build_solution(test, hangso,
                                os.path.join(OUTDIR, base + "_Loi_giai.docx"))
            made += [p1, p2]
            print("  ✓ %-26s  %2d + %d + %d câu"
                  % (base, len(test["p1"]), len(test["p2"]), len(test["p3"])))
    print("\nĐã tạo %d tệp .docx trong thư mục %s" % (len(made), os.path.basename(OUTDIR)))
    print("Phân bố đáp án Phần I trên toàn bộ 10 đề:", counts)
    return made


if __name__ == "__main__":
    main()
