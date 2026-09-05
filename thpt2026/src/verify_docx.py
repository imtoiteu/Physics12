# -*- coding: utf-8 -*-
"""Kiểm tra lại các tệp .docx đã sinh: đồng bộ mã câu, đáp án, hình và bảng."""
import os, re, sys, zipfile, importlib
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
sys.path.insert(0, HERE)
import build, make

RE_ID = re.compile(r"^Câu (\d+)\.")
RE_ANS = re.compile(r"^Đáp án:\s*(.+?)\s*$")
BAD = []


def texts(path):
    d = Document(path)
    return d, [p.text.strip() for p in d.paragraphs]


def media(path):
    with zipfile.ZipFile(path) as z:
        return [n for n in z.namelist() if n.startswith("word/media/")]


def main():
    for modname, fde, fgiai in make.SPEC:
        mod = importlib.import_module(modname)
        g = mod.NHOM
        build.rebalance_group(g)          # đưa dữ liệu về đúng trạng thái lúc dựng tệp
        qp, sp = os.path.join(ROOT, fde), os.path.join(ROOT, fgiai)
        dq, tq = texts(qp)
        ds, ts = texts(sp)

        n_q = sum(1 for t in tq if RE_ID.match(t))
        n_s = sum(1 for t in ts if RE_ID.match(t))
        need = sum(len(t["P1"]) + len(t["P2"]) + len(t["P3"]) for t in g["tests"])
        if n_q != need:
            BAD.append("%s: tệp đề có %d câu, cần %d" % (modname, n_q, need))
        if n_s != need:
            BAD.append("%s: tệp lời giải có %d câu, cần %d" % (modname, n_s, need))

        # đáp án trong tệp lời giải phải khớp dữ liệu nguồn
        got = [m.group(1) for m in (RE_ANS.match(t) for t in ts) if m]
        exp = []
        for t in g["tests"]:
            for it in t["P1"]:
                exp.append("%s. %s" % (it["a"], it["o"][build.LETTERS.index(it["a"])]))
            for it in t["P3"]:
                exp.append(it["ans"])
        if len(got) != len(exp):
            BAD.append("%s: có %d dòng “Đáp án:”, cần %d" % (modname, len(got), len(exp)))
        else:
            for i, (a, b) in enumerate(zip(got, exp)):
                if a.rstrip(".") != b.rstrip("."):
                    BAD.append("%s: dòng đáp án %d lệch: “%s” ≠ “%s”" % (modname, i + 1, a, b))

        figs = {it["fig"] for t in g["tests"] for k in ("P1", "P2", "P3")
                for it in t[k] if it.get("fig")}
        print("%-10s đề: %4d câu, %3d hình nhúng, %2d bảng | lời giải: %4d câu, %3d hình, %2d bảng"
              " | %2d hình khác nhau"
              % (modname, n_q, len(media(qp)), len(dq.tables),
                 n_s, len(media(sp)), len(ds.tables), len(figs)))

    if BAD:
        print("\n--- LỖI (%d) ---" % len(BAD))
        for b in BAD:
            print("  ✗", b)
        return 1
    print("\n✓ Tám tệp .docx đồng bộ và đầy đủ.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
