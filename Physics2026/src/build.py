# -*- coding: utf-8 -*-
"""Sinh bộ tài liệu Word cho NGÂN HÀNG ĐỀ LUYỆN THI TỐT NGHIỆP THPT 2026 – MÔN VẬT LÍ."""
import re
import os
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
FIGS = os.path.join(ROOT, "figs")
sys.path.insert(0, HERE)

FONT = "Times New Roman"
LETTERS = ["A", "B", "C", "D"]
DS_LABELS = ["a)", "b)", "c)", "d)"]

DARKRED = RGBColor(0xA6, 0x1B, 0x1B)
DARKBLUE = RGBColor(0x1F, 0x3A, 0x6E)
DARKGREEN = RGBColor(0x1E, 0x64, 0x36)
GREYTXT = RGBColor(0x55, 0x55, 0x55)
MDCOLOR = {"Dễ": RGBColor(0x1E, 0x64, 0x36), "Trung bình": RGBColor(0x1F, 0x3A, 0x6E),
           "Khó": RGBColor(0xD3, 0x54, 0x00), "Rất khó": RGBColor(0xA6, 0x1B, 0x1B)}
MUC_LIST = ["Dễ", "Trung bình", "Khó", "Rất khó"]


# ------------------------------------------------------------------ tiện ích
def set_base(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(11.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.line_spacing = 1.10
    for s in doc.sections:
        s.top_margin = Cm(1.7); s.bottom_margin = Cm(1.7)
        s.left_margin = Cm(2.0); s.right_margin = Cm(1.6)


def P(doc, text="", bold=False, italic=False, size=11.5, align=None, before=0, after=2,
      indent=None, color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if indent is not None:
        pf.left_indent = Cm(indent)
    if text:
        for i, line in enumerate(str(text).split("\n")):
            if i:
                p.add_run().add_break()
            r = p.add_run(line)
            r.bold = bold; r.italic = italic
            r.font.size = Pt(size); r.font.name = FONT
            if color is not None:
                r.font.color.rgb = color
    return p


def rich(doc, parts, indent=None, before=0, after=2, size=11.5, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if indent is not None:
        pf.left_indent = Cm(indent)
    for text, b, i, c in parts:
        for k, line in enumerate(str(text).split("\n")):
            if k:
                p.add_run().add_break()
            r = p.add_run(line)
            r.bold = b; r.italic = i; r.font.size = Pt(size); r.font.name = FONT
            if c is not None:
                r.font.color.rgb = c
    return p


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def box(doc, title, text, fill="EEF3FA", tcolor=DARKBLUE, size=11):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0); shade(c, fill)
    c.paragraphs[0].paragraph_format.space_after = Pt(2)
    if title:
        r = c.paragraphs[0].add_run(title)
        r.bold = True; r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = tcolor
    tp = c.add_paragraph(); tp.paragraph_format.space_after = Pt(1)
    for i, line in enumerate(str(text).split("\n")):
        if i:
            tp.add_run().add_break()
        r = tp.add_run(line); r.font.size = Pt(size); r.font.name = FONT
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def picture(doc, name, caption=None, maxw=11.5, maxh=7.4):
    path = os.path.join(FIGS, name + ".png")
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    with Image.open(path) as im:
        w, h = im.size
    width = maxw
    if h / w * width > maxh:
        width = maxh * w / h
    doc.add_picture(path, width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.space_before = Pt(3)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(1)
    if caption:
        P(doc, caption, italic=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER,
          after=4, color=GREYTXT)


def table(doc, caption, headers, rows, size=10):
    if caption:
        P(doc, caption, italic=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER,
          before=3, after=2, color=GREYTXT)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]; shade(c, "DCE6F1")
        c.paragraphs[0].paragraph_format.space_after = Pt(1)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.paragraphs[0].add_run(h); r.bold = True
        r.font.size = Pt(size); r.font.name = FONT
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].paragraphs[0].paragraph_format.space_after = Pt(1)
            if j:
                cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cells[j].paragraphs[0].add_run(str(v))
            r.font.size = Pt(size); r.font.name = FONT
            if j == 0:
                r.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_numbers(doc):
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(); run.font.size = Pt(9); run.font.name = FONT
        for instr in ("begin", "PAGE", "end"):
            if instr in ("begin", "end"):
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), instr)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve"); el.text = " PAGE "
            run._r.append(el)


# ------------------------------------------------------------------ nội dung
def kind_of(it):
    if "o" in it:
        return "mc"
    if "items" in it:
        return "ds"
    return "sa"


def short_answer(it):
    k = kind_of(it)
    if k == "mc":
        return it["a"]
    if k == "ds":
        return "".join("Đ" if v else "S" for _t, v, _e in it["items"])
    return it["ans"]


PART_HEAD = [
    ("I", "PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn",
     "Thí sinh trả lời từ câu 1 đến câu 18. Mỗi câu hỏi thí sinh chỉ chọn một phương án."),
    ("II", "PHẦN II. Câu trắc nghiệm đúng sai",
     "Thí sinh trả lời từ câu 1 đến câu 4. Trong mỗi ý a), b), c), d) ở mỗi câu, "
     "thí sinh chọn đúng hoặc sai."),
    ("III", "PHẦN III. Câu trắc nghiệm trả lời ngắn",
     "Thí sinh trả lời từ câu 1 đến câu 6."),
]


def parts_of(test):
    return [("I", test["P1"]), ("II", test["P2"]), ("III", test["P3"])]


def all_items(test):
    for code, items in parts_of(test):
        for i, it in enumerate(items, 1):
            yield code, i, it


def render_question(doc, num, it, part):
    rich(doc, [("Câu %d. " % num, True, False, None),
               (it["stem"] if part == "II" else it["q"], False, False, None)],
         before=5, after=2)
    if it.get("tbl"):
        table(doc, *it["tbl"])
    if it.get("fig"):
        picture(doc, it["fig"], it.get("cap"))
    if part == "II":
        for lab, (txt, _v, _e) in zip(DS_LABELS, it["items"]):
            rich(doc, [(lab + " ", True, False, None), (txt, False, False, None)],
                 indent=0.6, after=1)
    elif part == "I":
        opts = it["o"]
        short = all(len(o) <= 22 for o in opts)
        if short:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(1)
            for i, o in enumerate(opts):
                r = p.add_run("%s. " % LETTERS[i]); r.bold = True
                r.font.size = Pt(11.5); r.font.name = FONT
                r2 = p.add_run(o.ljust(24) if i < 3 else o)
                r2.font.size = Pt(11.5); r2.font.name = FONT
        else:
            for i, o in enumerate(opts):
                rich(doc, [("%s. " % LETTERS[i], True, False, None), (o, False, False, None)],
                     indent=0.6, after=0)
    else:
        P(doc, "Đáp số:  ..................................", size=11, indent=0.6, after=1)


def render_solution(doc, num, it, part):
    md = it.get("md", "Trung bình")
    rich(doc, [("Câu %d. " % num, True, False, DARKRED),
               ("[%s]  " % md, True, True, MDCOLOR.get(md, GREYTXT)),
               ("Kiến thức – kĩ năng: %s" % it.get("kn", ""), False, True, DARKGREEN)],
         before=6, after=1, size=10.5)
    rich(doc, [(it["stem"] if part == "II" else it["q"], False, True, None)], after=2)
    if it.get("tbl"):
        table(doc, *it["tbl"])
    if it.get("fig"):
        picture(doc, it["fig"], it.get("cap"), maxw=9.5, maxh=6.2)
    if part == "II":
        for lab, (txt, val, ex) in zip(DS_LABELS, it["items"]):
            rich(doc, [(lab + " ", True, False, None), (txt, False, False, None),
                       ("  →  " + ("ĐÚNG" if val else "SAI"), True, False,
                        DARKBLUE if val else DARKRED)], indent=0.6, after=1)
            rich(doc, [(ex, False, False, None)], indent=1.1, after=2, size=11)
        return
    if part == "I":
        rich(doc, [("Đáp án: %s. " % it["a"], True, False, DARKRED),
                   (it["o"][LETTERS.index(it["a"])], True, False, DARKRED)],
             indent=0.6, after=1)
    else:
        rich(doc, [("Đáp án: ", True, False, DARKRED), (it["ans"], True, False, DARKRED)],
             indent=0.6, after=1)
    rich(doc, [(it["sol"], False, False, None)], indent=0.6, after=2, size=11)


def answer_key(doc, test, ncol=9):
    P(doc, "BẢNG ĐÁP ÁN – %s" % test["ma"], bold=True, size=11.5,
      align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=3, color=DARKBLUE)
    for code, items in parts_of(test):
        pairs = [("%d" % i, short_answer(it)) for i, it in enumerate(items, 1)]
        P(doc, "Phần %s" % code, bold=True, size=10.5, after=2, color=DARKGREEN)
        n = min(ncol, len(pairs))
        nrow = (len(pairs) + n - 1) // n
        t = doc.add_table(rows=2 * nrow, cols=n)
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, (ident, ans) in enumerate(pairs):
            r0, c0 = 2 * (k // n), k % n
            for row, val, bold, col in ((r0, "Câu " + ident, True, None),
                                        (r0 + 1, ans, True, DARKRED)):
                cell = t.rows[row].cells[c0]
                if row == r0:
                    shade(cell, "F0F3F7")
                cell.paragraphs[0].paragraph_format.space_after = Pt(0)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                rr = cell.paragraphs[0].add_run(val)
                rr.bold = bold; rr.font.size = Pt(9.5); rr.font.name = FONT
                if col is not None:
                    rr.font.color.rgb = col
        doc.add_paragraph().paragraph_format.space_after = Pt(2)


def test_header(doc, test, group, solution=False):
    P(doc, group["ten_nhom"], bold=True, size=11,
      align=WD_ALIGN_PARAGRAPH.CENTER, after=0, color=GREYTXT)
    P(doc, "%s  –  %s" % (test["ma"], test["ten"]), bold=True, size=15,
      align=WD_ALIGN_PARAGRAPH.CENTER, after=1, color=DARKRED)
    P(doc, "Mức độ: %s   |   Thời gian làm bài: 50 phút   |   Thang điểm: 10"
      % test["muc"], italic=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    P(doc, "Trọng tâm: " + test["trongtam"], italic=True, size=10.5,
      align=WD_ALIGN_PARAGRAPH.CENTER, after=5, color=GREYTXT)


SCORING = ("Cấu trúc và cách tính điểm (theo đề tham khảo/chính thức của Bộ GD&ĐT áp dụng từ 2025):\n"
           "• Phần I – 18 câu trắc nghiệm nhiều phương án lựa chọn, mỗi câu 0,25 điểm → 4,50 điểm.\n"
           "• Phần II – 4 câu trắc nghiệm đúng/sai, mỗi câu 4 ý. Trong một câu: đúng 1 ý được 0,10 điểm; "
           "đúng 2 ý được 0,25 điểm; đúng 3 ý được 0,50 điểm; đúng cả 4 ý được 1,00 điểm → 4,00 điểm.\n"
           "• Phần III – 6 câu trắc nghiệm trả lời ngắn, mỗi câu 0,25 điểm → 1,50 điểm.\n"
           "Tổng: 28 câu hỏi / 40 lệnh hỏi / 10,00 điểm / 50 phút.")


def build_de(group, filename, outdir=None):
    doc = Document(); set_base(doc); add_page_numbers(doc)
    cover(doc, group, "TẬP ĐỀ BÀI")
    box(doc, "CẤU TRÚC ĐỀ VÀ CÁCH TÍNH ĐIỂM", SCORING, fill="F2F8F0", tcolor=DARKGREEN)
    matrix_table(doc, group)
    doc.add_page_break()
    for ti, test in enumerate(group["tests"]):
        test_header(doc, test, group)
        for code, items in parts_of(test):
            _c, head, note = [p for p in PART_HEAD if p[0] == code][0]
            P(doc, head, bold=True, size=12, before=8, after=1, color=DARKBLUE)
            P(doc, note, italic=True, size=10, after=3, color=GREYTXT)
            for i, it in enumerate(items, 1):
                render_question(doc, i, it, code)
        P(doc, "----- HẾT %s -----" % test["ma"], bold=True, size=11,
          align=WD_ALIGN_PARAGRAPH.CENTER, before=8)
        if ti < len(group["tests"]) - 1:
            doc.add_page_break()
    path = os.path.join(outdir or ROOT, filename); doc.save(path)
    return path


def build_giai(group, filename, outdir=None):
    doc = Document(); set_base(doc); add_page_numbers(doc)
    cover(doc, group, "ĐÁP ÁN VÀ LỜI GIẢI CHI TIẾT")
    box(doc, "CÁCH ĐỌC TÀI LIỆU NÀY",
        "• Số hiệu đề và số thứ tự câu trùng khớp tuyệt đối với tập đề bài.\n"
        "• Mỗi câu đều ghi rõ MỨC ĐỘ (Dễ / Trung bình / Khó / Rất khó) và KIẾN THỨC – KĨ NĂNG được kiểm tra.\n"
        "• Với câu đúng/sai, mỗi ý đều được giải thích riêng vì sao đúng hoặc vì sao sai.\n"
        "• Bảng đáp án nhanh của mỗi đề được đặt ngay trước phần lời giải của đề đó.",
        fill="F4F6F8")
    doc.add_page_break()
    for ti, test in enumerate(group["tests"]):
        test_header(doc, test, group, solution=True)
        answer_key(doc, test)
        stat_table(doc, test)
        for code, items in parts_of(test):
            _c, head, _note = [p for p in PART_HEAD if p[0] == code][0]
            P(doc, "LỜI GIẢI – " + head, bold=True, size=12, before=8, after=3,
              color=DARKBLUE)
            for i, it in enumerate(items, 1):
                render_solution(doc, i, it, code)
        if ti < len(group["tests"]) - 1:
            doc.add_page_break()
    path = os.path.join(outdir or ROOT, filename); doc.save(path)
    return path


def stat_table(doc, test):
    cnt = {m: 0 for m in MUC_LIST}
    for _c, _i, it in all_items(test):
        cnt[it.get("md", "Trung bình")] = cnt.get(it.get("md", "Trung bình"), 0) + 1
    tot = sum(cnt.values())
    rows = [["Số câu"] + [str(cnt[m]) for m in MUC_LIST] + [str(tot)],
            ["Tỉ lệ"] + ["%.0f%%" % (100 * cnt[m] / tot) for m in MUC_LIST] + ["100%"]]
    table(doc, "Phân bố mức độ của %s" % test["ma"], ["Mức độ"] + MUC_LIST + ["Tổng"], rows)


def matrix_table(doc, group):
    n = len(group["tests"])
    tieu_de = "MA TRẬN ĐỘ KHÓ CỦA ĐỀ" if n == 1 else "MA TRẬN ĐỘ KHÓ CỦA BỘ %d ĐỀ" % n
    P(doc, tieu_de, bold=True, size=12.5,
      align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=3, color=DARKBLUE)
    rows = []
    for t in group["tests"]:
        cnt = {m: 0 for m in MUC_LIST}
        for _c, _i, it in all_items(t):
            cnt[it.get("md", "Trung bình")] += 1
        rows.append([t["ma"], t["muc"]] + [str(cnt[m]) for m in MUC_LIST] + [t["trongtam"]])
    table(doc, None, ["Mã đề", "Mức độ chung"] + MUC_LIST + ["Trọng tâm"], rows, size=9)


def cover(doc, group, kind_label):
    P(doc, "NGÂN HÀNG ĐỀ LUYỆN THI TỐT NGHIỆP THPT 2026", bold=True, size=12,
      align=WD_ALIGN_PARAGRAPH.CENTER, after=0, color=GREYTXT)
    P(doc, "MÔN VẬT LÍ", bold=True, size=26, align=WD_ALIGN_PARAGRAPH.CENTER,
      after=1, color=DARKRED)
    P(doc, "Biên soạn theo Chương trình GDPT 2018 và cấu trúc định dạng đề thi "
           "tốt nghiệp THPT từ năm 2025", italic=True, size=11,
      align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    P(doc, group["ten_nhom"], bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER,
      after=2, color=DARKBLUE)
    P(doc, kind_label, bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    P(doc, group["mo_ta"], italic=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER,
      after=10)
    box(doc, "PHẠM VI KIẾN THỨC", group["pham_vi"], fill="F4F6F8")


# ------------------------------------------------------------ cân bằng đáp án
def _numeric_options(opts):
    n = 0
    for o in opts:
        t = o.strip().lstrip("≈±~").lstrip()
        if t[:1].isdigit() or t[:1] in "-−":
            n += 1
    return n >= 3


SUP = {"⁻": "-", "⁰": "0", "¹": "1", "²": "2", "³": "3", "⁴": "4",
       "⁵": "5", "⁶": "6", "⁷": "7", "⁸": "8", "⁹": "9"}
_NUM = re.compile(r"[-−]?\d+(?:[,.]\d+)?")


def _optval(o):
    """Giá trị số của một phương án, dùng để biết bộ phương án có sắp thứ tự hay không."""
    t = o.strip().lstrip("≈±~").lstrip()
    m = _NUM.match(t)
    if not m:
        return None
    try:
        v = float(m.group(0).replace("−", "-").replace(",", "."))
    except ValueError:
        return None
    rest = t[m.end():]
    e = re.match(r"\s*[·x×]\s*10([⁻⁰¹²³⁴⁵⁶⁷⁸⁹]+)", rest)
    if e:
        try:
            v *= 10.0 ** int("".join(SUP.get(c, "") for c in e.group(1)))
        except ValueError:
            return None
    return v


def _monotonic(opts):
    """True nếu bộ phương án là dãy số đã sắp tăng hoặc giảm — khi đó chỉ được phép đảo chiều."""
    vs = [_optval(o) for o in opts]
    if any(v is None for v in vs):
        return False
    up = all(vs[i] < vs[i + 1] for i in range(len(vs) - 1))
    dn = all(vs[i] > vs[i + 1] for i in range(len(vs) - 1))
    return up or dn


def rebalance(items):
    """Phân bố lại vị trí đáp án đúng cho đều giữa A, B, C, D trong MỘT đề.

    • Bộ phương án là giá trị số: chỉ cho phép đảo chiều sắp xếp (tăng ↔ giảm dần),
      phép đảo đưa đáp án A↔D và B↔C, vẫn giữ được thứ tự dễ đọc.
    • Bộ phương án dạng chữ: hoán vị tự do vị trí phương án đúng.
    Nội dung phương án không đổi nên lời giải vẫn đúng nguyên vẹn.
    """
    numeric = [it for it in items
               if _numeric_options(it["o"]) and _monotonic(it["o"])]
    pool = [it for it in items if it not in numeric]

    fixed = {L: 0 for L in LETTERS}
    for it in numeric:
        fixed[it["a"]] += 1
    for _ in range(5):
        changed = False
        for it in numeric:
            cur = it["a"]
            mir = LETTERS[3 - LETTERS.index(cur)]
            if fixed[cur] > fixed[mir] + 1:
                it["o"].reverse()
                it["a"] = mir
                fixed[cur] -= 1
                fixed[mir] += 1
                changed = True
        if not changed:
            break

    alloc = {L: 0 for L in LETTERS}
    for _ in pool:
        L = min(LETTERS, key=lambda x: (fixed[x] + alloc[x], LETTERS.index(x)))
        alloc[L] += 1
    slots = []
    for L in LETTERS:
        slots += [L] * alloc[L]
    need = []
    for it in pool:
        if it["a"] in slots:
            slots.remove(it["a"])
        else:
            need.append(it)
    for it in need:
        dst = slots.pop(0)
        i, j = LETTERS.index(it["a"]), LETTERS.index(dst)
        it["o"][i], it["o"][j] = it["o"][j], it["o"][i]
        it["a"] = dst

    out = {L: 0 for L in LETTERS}
    for it in items:
        out[it["a"]] += 1
    return out


def rebalance_group(group):
    tot = {L: 0 for L in LETTERS}
    for t in group["tests"]:
        c = rebalance(t["P1"])
        for L in LETTERS:
            tot[L] += c[L]
    return tot
