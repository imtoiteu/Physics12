# -*- coding: utf-8 -*-
"""Engine dựng slide bài giảng PowerPoint (.pptx) cho bộ tài liệu ôn thi THPT 2026.

Nguyên tắc trình bày
--------------------
* Khổ 16:9, bố cục cố định: dải tiêu đề – vạch màu nhấn – vùng nội dung – chân trang.
* Chữ dùng Arial (có sẵn trên mọi máy Windows/Office, đủ dấu tiếng Việt).
* Kí hiệu đơn giản (Δ, λ, μ, ω, °C, ·10³, m³) viết thẳng bằng Unicode.
* Công thức hai chiều (phân số, căn, gạch ngang trên) dựng thành ẢNH nét cao bằng
  bộ chữ toán STIX ⇒ tuyệt đối không lỗi font, không lộ mã LaTeX.
* Mọi khối chữ đều được ĐO trước bằng phông Liberation Sans (cùng hệ số với Arial);
  cỡ chữ tự co cho vừa khung, nếu vẫn tràn thì báo lỗi để tác giả tách slide.
"""
import os
import sys

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
FIGS = os.path.join(ROOT, "figs")
sys.path.insert(0, HERE)

import formula                                                    # noqa: E402

FONT = "Arial"
SW, SH = 13.3333, 7.5

_METRIC = "/usr/share/fonts/truetype/liberation/LiberationSans-%s.ttf"
_FCACHE = {}


def _font(pt, bold=False):
    key = (round(pt * 2), bold)
    if key not in _FCACHE:
        path = _METRIC % ("Bold" if bold else "Regular")
        _FCACHE[key] = ImageFont.truetype(path, int(round(pt * 4)))
    return _FCACHE[key]


def textw(s, pt, bold=False):
    """Bề rộng chuỗi tính theo inch khi in ở cỡ pt."""
    f = _font(pt, bold)
    return f.getlength(s) / 4.0 / 72.0


def wrapn(s, pt, w_in, bold=False):
    """Số dòng mà chuỗi s chiếm khi bọc trong khung rộng w_in inch."""
    if not s:
        return 1
    words, lines, cur = s.split(), 0, ""
    for wd in words:
        trial = (cur + " " + wd).strip()
        if textw(trial, pt, bold) <= w_in or not cur:
            cur = trial
        else:
            lines += 1
            cur = wd
    return lines + (1 if cur else 0)


# ----------------------------------------------------------------- bảng màu
def theme(primary, dark, tint, soft):
    return dict(primary=RGBColor.from_string(primary), dark=RGBColor.from_string(dark),
                tint=tint, soft=soft, primary_s=primary, dark_s=dark)


THEMES = {
    1: theme("C0392B", "7B241C", "FDEDEC", "FBE3E0"),   # Nhiệt  – đỏ cam
    2: theme("117A65", "0B4F42", "E8F6F3", "D6EDE8"),   # Khí    – xanh ngọc
    3: theme("2E4A9E", "1B2E63", "EAEFFA", "DCE4F6"),   # Từ     – xanh chàm
    4: theme("6C3483", "4A235A", "F4ECF7", "EADCF0"),   # Hạt nhân – tím
    0: theme("34495E", "1C2833", "EDF1F4", "E0E6EA"),   # trung tính
}

INK = RGBColor.from_string("1A1A1A")
INK2 = RGBColor.from_string("3D3D3D")
GREY = RGBColor.from_string("7A7A7A")
WHITE = RGBColor.from_string("FFFFFF")
GOLD = RGBColor.from_string("B8860B")

# vùng nội dung chuẩn
BX, BY, BW, BH = 0.70, 1.52, 11.95, 5.30


class Deck(object):
    def __init__(self, title, chapter, subtitle="", meta="", th=None):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SW)
        self.prs.slide_height = Inches(SH)
        self.blank = self.prs.slide_layouts[6]
        self.th = THEMES[th if th is not None else 0]
        self.title = title
        self.chapter = chapter
        self.subtitle = subtitle
        self.meta = meta
        self.n = 0
        self.warn = []

    # ------------------------------------------------------------ nguyên thuỷ
    def _slide(self):
        return self.prs.slides.add_slide(self.blank)

    def _rect(self, s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, adj=None):
        sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.shadow.inherit = False
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid()
            sh.fill.fore_color.rgb = fill if isinstance(fill, RGBColor) else RGBColor.from_string(fill)
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line if isinstance(line, RGBColor) else RGBColor.from_string(line)
            sh.line.width = Pt(1.1)
        if adj is not None:
            try:
                sh.adjustments[0] = adj
            except Exception:
                pass
        sh.text_frame.text = ""
        return sh

    def _tb(self, s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Emu(0)
        tf.margin_top = tf.margin_bottom = Emu(0)
        return tf

    @staticmethod
    def _run(p, text, pt, bold=False, italic=False, color=INK, font=FONT):
        r = p.add_run()
        r.text = text
        r.font.size = Pt(pt)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
        return r

    def _line(self, tf, text, pt, bold=False, italic=False, color=INK, align=PP_ALIGN.LEFT,
              space_before=0, space_after=0, first=False, indent=0.0):
        p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(space_before)
        p.space_after = Pt(space_after)
        if indent:
            p.left_indent = Inches(indent)
        if isinstance(text, (list, tuple)):
            for seg in text:
                t, b, i, c = (seg + (None,) * 4)[:4]
                self._run(p, t, pt, b or False, i or False, c or color)
        else:
            self._run(p, text, pt, bold, italic, color)
        return p

    def _pic(self, s, path, x, y, w=None, h=None):
        kw = {}
        if w:
            kw["width"] = Inches(w)
        if h:
            kw["height"] = Inches(h)
        return s.shapes.add_picture(path, Inches(x), Inches(y), **kw)

    def _fx(self, s, tex, cx, y, maxh=0.52, maxw=11.0, color=None):
        """Đặt ảnh công thức, canh giữa theo trục cx.  Trả về chiều cao thực tế."""
        col = color or ("#%s" % self.th["dark_s"])
        path, w, h = formula.render(tex, color=col)
        sc = maxh / h
        if w * sc > maxw:
            sc = maxw / w
        w, h = w * sc, h * sc
        self._pic(s, path, cx - w / 2.0, y, w=w)
        return h

    # ------------------------------------------------------------ khung slide
    def _chrome(self, s, head, sub=None, tag=None):
        th = self.th
        self._rect(s, 0, 0, SW, 0.13, th["primary"])
        tf = self._tb(s, BX, 0.40, 11.2, 0.92)
        self._line(tf, head, 27 if len(head) <= 52 else 23, bold=True, color=th["dark"],
                   first=True)
        if sub:
            self._line(tf, sub, 13, italic=True, color=GREY, space_before=2)
        self._rect(s, BX, 1.33 if not sub else 1.36, 1.45, 0.05, th["primary"])
        if tag:
            w = max(1.1, textw(tag, 11, True) + 0.42)
            self._rect(s, SW - 0.70 - w, 0.46, w, 0.36, th["tint"], shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.35)
            tf2 = self._tb(s, SW - 0.70 - w, 0.52, w, 0.3, MSO_ANCHOR.MIDDLE)
            self._line(tf2, tag, 11, bold=True, color=th["dark"], align=PP_ALIGN.CENTER, first=True)
        self.n += 1
        ftf = self._tb(s, BX, 6.98, 9.6, 0.3)
        self._line(ftf, self.chapter, 9, color=GREY, first=True)
        ntf = self._tb(s, SW - 1.55, 6.98, 0.85, 0.3)
        self._line(ntf, str(self.n), 9, bold=True, color=self.th["primary"],
                   align=PP_ALIGN.RIGHT, first=True)

    @staticmethod
    def _notes(s, text):
        if text:
            s.notes_slide.notes_text_frame.text = text

    # ------------------------------------------------------------ các layout
    def title_slide(self, sub=None, meta=None, points=()):
        s = self._slide()
        th = self.th
        self._rect(s, 0, 0, 0.42, SH, th["primary"])
        self._rect(s, 0.42, 0, SW - 0.42, SH, th["tint"])
        nline = wrapn(self.title, 38 if len(self.title) <= 46 else 31, 10.1, True)
        cardh = 3.30 + 0.62 * (nline - 1)
        self._rect(s, 1.05, 1.05, 10.9, cardh, WHITE)
        tf = self._tb(s, 1.45, 1.35, 10.1, 0.4)
        self._line(tf, "VẬT LÍ 12  •  ÔN THI TỐT NGHIỆP THPT 2026", 12.5, bold=True,
                   color=th["primary"], first=True)
        tf2 = self._tb(s, 1.45, 1.92, 10.1, 0.72 * nline + 0.9)
        self._line(tf2, self.title, 38 if len(self.title) <= 46 else 31, bold=True,
                   color=th["dark"], first=True, space_after=6)
        if sub or self.subtitle:
            self._line(tf2, sub or self.subtitle, 17, color=INK2, space_before=4)
        self._rect(s, 1.45, 1.05 + cardh - 0.92, 2.0, 0.06, th["primary"])
        tf3 = self._tb(s, 1.45, 1.05 + cardh - 0.70, 10.1, 0.45)
        self._line(tf3, meta or self.meta, 12.5, italic=True, color=GREY, first=True)
        if points:
            tf4 = self._tb(s, 1.05, 1.05 + cardh + 0.34, 11.0, 1.5)
            for i, p in enumerate(points):
                self._line(tf4, [("▪  ", True, False, th["primary"]), (p, False, False, INK2)],
                           12.5, first=(i == 0), space_after=3)
        return s

    def section(self, num, head, desc="", items=()):
        s = self._slide()
        th = self.th
        self._rect(s, 0, 0, SW, SH, th["dark"])
        self._rect(s, 0, SH - 0.16, SW, 0.16, th["primary"])
        tf = self._tb(s, 1.15, 2.05, 3.0, 1.5)
        self._line(tf, "%02d" % num, 76, bold=True, color=th["primary"], first=True)
        tf2 = self._tb(s, 3.15, 2.15, 9.2, 2.6)
        self._line(tf2, head, 33 if len(head) <= 40 else 27, bold=True, color=WHITE, first=True)
        if desc:
            self._line(tf2, desc, 14.5, color=RGBColor.from_string("D5DBE0"), space_before=10)
        if items:
            tf3 = self._tb(s, 3.20, 4.55, 9.0, 1.9)
            for i, it in enumerate(items):
                self._line(tf3, [("—  ", False, False, th["primary"]),
                                 (it, False, False, RGBColor.from_string("E8ECEF"))],
                           13, first=(i == 0), space_after=4)
        self.n += 1
        return s

    # -------- thân slide: hệ phần tử
    def _emit(self, s, x, y, w, items, pt, tight=False):
        """Vẽ danh sách phần tử nội dung.  Trả về y sau cùng."""
        th = self.th
        gap = 0.055 if tight else 0.085
        for it in items:
            k, payload = _norm(it)
            if k == "fx":
                y += 0.05
                y += self._fx(s, payload, x + w / 2.0, y, maxh=min(0.60, pt / 46.0), maxw=w * 0.95)
                y += 0.10
            elif k == "gap":
                y += float(payload)
            elif k == "rule":
                self._rect(s, x, y + 0.06, w, 0.014, th["soft"])
                y += 0.16
            elif k == "head":
                h = wrapn(payload, pt + 1.5, w, True) * (pt + 1.5) / 72.0 * 1.28
                tf = self._tb(s, x, y, w, h + 0.1)
                self._line(tf, payload, pt + 1.5, bold=True, color=th["dark"], first=True)
                y += h + gap
            elif k == "note":
                h = wrapn(payload, pt - 1.5, w - 0.35) * (pt - 1.5) / 72.0 * 1.3
                self._rect(s, x, y, 0.045, h + 0.10, th["primary"])
                tf = self._tb(s, x + 0.22, y + 0.03, w - 0.25, h + 0.1)
                self._line(tf, payload, pt - 1.5, italic=True, color=INK2, first=True)
                y += h + 0.16
            elif k in ("b", "sub", "num"):
                mark, ind, size = {"b": ("▪", 0.0, pt), "sub": ("–", 0.34, pt - 1.0),
                                   "num": ("", 0.0, pt)}[k]
                if k == "num":
                    mark, payload = payload
                lead = (mark + "  ") if mark else ""
                lw = textw(lead, size, True)
                h = wrapn(payload, size, w - ind - lw) * size / 72.0 * 1.30
                tfm = self._tb(s, x + ind, y, lw + 0.05, h)
                self._line(tfm, mark, size, bold=True, color=th["primary"], first=True)
                tf = self._tb(s, x + ind + lw, y, w - ind - lw, h + 0.05)
                self._line(tf, payload, size, color=INK, first=True)
                y += h + gap
            elif k == "kv":
                key, val = payload
                kw_ = min(3.6, textw(key, pt, True) + 0.15)
                h = max(wrapn(val, pt, w - kw_ - 0.2), 1) * pt / 72.0 * 1.30
                tf = self._tb(s, x, y, kw_, h)
                self._line(tf, key, pt, bold=True, color=th["primary"], first=True)
                tf2 = self._tb(s, x + kw_ + 0.15, y, w - kw_ - 0.15, h + 0.05)
                self._line(tf2, val, pt, color=INK, first=True)
                y += h + gap
            else:                                   # 'p' – đoạn văn thường
                h = wrapn(payload, pt, w) * pt / 72.0 * 1.30
                tf = self._tb(s, x, y, w, h + 0.05)
                self._line(tf, payload, pt, color=INK, first=True)
                y += h + gap
        return y

    def _need(self, items, pt, w):
        """Chiều cao ước lượng của danh sách phần tử."""
        tot = 0.0
        for it in items:
            k, payload = _norm(it)
            if k == "fx":
                tot += min(0.60, pt / 46.0) + 0.15
            elif k == "gap":
                tot += float(payload)
            elif k == "rule":
                tot += 0.16
            elif k == "head":
                tot += wrapn(payload, pt + 1.5, w, True) * (pt + 1.5) / 72.0 * 1.28 + 0.085
            elif k == "note":
                tot += wrapn(payload, pt - 1.5, w - 0.35) * (pt - 1.5) / 72.0 * 1.3 + 0.16
            elif k in ("b", "sub", "num"):
                ind, size = {"b": (0.0, pt), "sub": (0.34, pt - 1.0), "num": (0.0, pt)}[k]
                if k == "num":
                    payload = payload[1]
                tot += wrapn(payload, size, w - ind - 0.30) * size / 72.0 * 1.30 + 0.085
            elif k == "kv":
                key, val = payload
                kw_ = min(3.6, textw(key, pt, True) + 0.15)
                tot += wrapn(val, pt, w - kw_ - 0.2) * pt / 72.0 * 1.30 + 0.085
            else:
                tot += wrapn(payload, pt, w) * pt / 72.0 * 1.30 + 0.085
        return tot

    def _autofit(self, items, w, h, base=17.0, low=11.5, high=None):
        """Chọn cỡ chữ lớn nhất mà nội dung vẫn vừa khung (co xuống hoặc giãn lên)."""
        high = high if high is not None else base + 3.0
        pt = base
        if self._need(items, pt, w) <= h:
            while pt + 0.5 <= high and self._need(items, pt + 0.5, w) <= h:
                pt += 0.5
        else:
            while pt > low and self._need(items, pt, w) > h:
                pt -= 0.5
        over = self._need(items, pt, w) - h
        return pt, over

    def _voff(self, items, pt, w, h, mode="auto"):
        """Độ lệch dọc để khối nội dung ngắn không dồn hết lên đỉnh slide."""
        need = self._need(items, pt, w)
        if mode == "top" or need >= h:
            return 0.0
        if mode == "center":
            return (h - need) / 2.0
        return min((h - need) / 2.0, (h - need) * 0.42)

    # -------- các loại slide nội dung
    def content(self, head, items, sub=None, tag=None, note=None, base=17.0):
        s = self._slide()
        self._chrome(s, head, sub, tag)
        y0 = BY + (0.06 if sub else 0)
        H = BH - (0.06 if sub else 0)
        pt, over = self._autofit(items, BW, H, base=base)
        if over > 0:
            self.warn.append("TRÀN %.2f in — slide %d «%s»" % (over, self.n, head))
        self._emit(s, BX, y0 + self._voff(items, pt, BW, H), BW, items, pt)
        self._notes(s, note)
        return s

    def split(self, head, left, right, sub=None, tag=None, note=None, ratio=0.5,
              lhead=None, rhead=None, base=16.0):
        s = self._slide()
        self._chrome(s, head, sub, tag)
        gapx = 0.42
        lw = (BW - gapx) * ratio
        rw = BW - gapx - lw
        y0 = BY + (0.06 if sub else 0)
        H = BH - (0.06 if sub else 0)
        if lhead:
            left = [("head", lhead)] + list(left)
        if rhead:
            right = [("head", rhead)] + list(right)
        pt = min(self._autofit(left, lw, H, base=base)[0],
                 self._autofit(right, rw, H, base=base)[0])
        for items, w in ((left, lw), (right, rw)):
            if self._need(items, pt, w) > H:
                self.warn.append("TRÀN cột — slide %d «%s»" % (self.n, head))
        self._rect(s, BX + lw + gapx / 2 - 0.007, y0 + 0.05, 0.014, H - 0.2, self.th["soft"])
        self._emit(s, BX, y0 + self._voff(left, pt, lw, H), lw, left, pt)
        self._emit(s, BX + lw + gapx, y0 + self._voff(right, pt, rw, H), rw, right, pt)
        self._notes(s, note)
        return s

    def figure(self, head, fig, items=(), cap=None, sub=None, tag=None, note=None,
               figw=6.1, side="right", base=16.0):
        s = self._slide()
        self._chrome(s, head, sub, tag)
        path = os.path.join(FIGS, fig + ".png")
        if not os.path.exists(path):
            raise FileNotFoundError(path)
        with Image.open(path) as im:
            iw, ih = im.size
        y0 = BY + (0.06 if sub else 0)
        H = BH - (0.06 if sub else 0)
        w = figw
        h = w * ih / iw
        if h > H - (0.30 if cap else 0.05):
            h = H - (0.30 if cap else 0.05)
            w = h * iw / ih
        tw = BW - figw - 0.45
        fx = BX + BW - figw + (figw - w) / 2.0 if side == "right" else BX + (figw - w) / 2.0
        tx = BX if side == "right" else BX + figw + 0.45
        self._rect(s, fx - 0.10, y0 + (H - h) / 2.0 - 0.10, w + 0.20, h + 0.20, "FFFFFF",
                   line=self.th["soft"])
        self._pic(s, path, fx, y0 + (H - h) / 2.0, w=w)
        if cap:
            tfc = self._tb(s, fx - 0.3, y0 + (H - h) / 2.0 + h + 0.13, w + 0.6, 0.3)
            self._line(tfc, cap, 10.5, italic=True, color=GREY, align=PP_ALIGN.CENTER, first=True)
        if items:
            pt, over = self._autofit(items, tw, H, base=base)
            if over > 0:
                self.warn.append("TRÀN %.2f in — slide %d «%s»" % (over, self.n, head))
            self._emit(s, tx, y0 + self._voff(items, pt, tw, H, "center"), tw, items, pt)
        self._notes(s, note)
        return s

    def bigfigure(self, head, fig, cap=None, sub=None, tag=None, note=None, items=()):
        s = self._slide()
        self._chrome(s, head, sub, tag)
        path = os.path.join(FIGS, fig + ".png")
        with Image.open(path) as im:
            iw, ih = im.size
        y0 = BY
        H = BH - (0.32 if cap else 0.05)
        bot = 0.0
        if items:
            bot = self._need(items, 14.0, BW) + 0.12
            H -= bot
        w, h = BW * 0.82, BW * 0.82 * ih / iw
        if h > H:
            h = H
            w = h * iw / ih
        x = BX + (BW - w) / 2.0
        self._rect(s, x - 0.10, y0 - 0.02, w + 0.20, h + 0.16, "FFFFFF", line=self.th["soft"])
        self._pic(s, path, x, y0 + 0.05, w=w)
        if cap:
            tfc = self._tb(s, BX, y0 + h + 0.22, BW, 0.3)
            self._line(tfc, cap, 11, italic=True, color=GREY, align=PP_ALIGN.CENTER, first=True)
        if items:
            self._emit(s, BX, y0 + h + (0.55 if cap else 0.28), BW, items, 14.0)
        self._notes(s, note)
        return s

    def formulas(self, head, rows, sub=None, tag=None, note=None, lead=None, foot=None):
        """rows: [(tex, chú thích)] — mỗi công thức một thẻ ngang."""
        s = self._slide()
        self._chrome(s, head, sub, tag)
        th = self.th
        y = BY + (0.06 if sub else 0)
        H = BH - (0.06 if sub else 0)
        if lead:
            hh = wrapn(lead, 15, BW) * 15 / 72.0 * 1.3
            tf = self._tb(s, BX, y, BW, hh + 0.05)
            self._line(tf, lead, 15, color=INK2, first=True)
            y += hh + 0.16
            H -= hh + 0.16
        fh = wrapn(foot, 13.5, BW) * 13.5 / 72.0 * 1.3 + 0.2 if foot else 0.0
        H -= fh
        n = len(rows)
        cardh = min(1.12, (H - 0.12 * (n - 1)) / n)
        fxh = min(0.60, cardh * 0.50)
        y += max(0.0, (H - (cardh * n + 0.12 * (n - 1))) / 2.0)
        for tex, cmt in rows:
            self._rect(s, BX, y, BW, cardh, th["tint"], shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.09)
            self._rect(s, BX, y, 0.075, cardh, th["primary"])
            cw = 4.9
            self._fx(s, tex, BX + 0.17 + cw / 2.0, y + (cardh - fxh) / 2.0,
                     maxh=fxh, maxw=cw - 0.3)
            if cmt:
                pt = 14.0
                while pt > 10.5 and wrapn(cmt, pt, BW - cw - 0.55) * pt / 72.0 * 1.3 > cardh - 0.18:
                    pt -= 0.5
                hh = wrapn(cmt, pt, BW - cw - 0.55) * pt / 72.0 * 1.3
                tf = self._tb(s, BX + cw + 0.35, y + max(0.09, (cardh - hh) / 2.0),
                              BW - cw - 0.55, hh + 0.06)
                self._line(tf, cmt, pt, color=INK, first=True)
            y += cardh + 0.12
        if foot:
            self._rect(s, BX, y + 0.02, 0.045, fh - 0.1, th["primary"])
            tf = self._tb(s, BX + 0.22, y + 0.04, BW - 0.25, fh)
            self._line(tf, foot, 13.5, italic=True, color=INK2, first=True)
        self._notes(s, note)
        return s

    def example(self, head, q, steps, ans, sub=None, tag="VÍ DỤ", note=None, fig=None,
                cap=None, tip=None):
        s = self._slide()
        self._chrome(s, head, sub, tag)
        th = self.th
        y = BY
        qh = wrapn(q, 15, BW - 0.45) * 15 / 72.0 * 1.32
        self._rect(s, BX, y, BW, qh + 0.30, th["tint"], shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.10)
        tf = self._tb(s, BX + 0.22, y + 0.15, BW - 0.45, qh + 0.05)
        self._line(tf, q, 15, color=INK, first=True)
        y += qh + 0.44
        H = BH - (y - BY) - (0.0 if not tip else 0.0)
        if fig:
            path = os.path.join(FIGS, fig + ".png")
            with Image.open(path) as im:
                iw, ih = im.size
            fw = 4.5
            fh = fw * ih / iw
            if fh > H - 0.1:
                fh = H - 0.1
                fw = fh * iw / ih
            self._pic(s, path, BX + BW - fw, y + (H - fh) / 2.0, w=fw)
            if cap:
                tfc = self._tb(s, BX + BW - fw - 0.2, y + (H - fh) / 2.0 + fh + 0.06, fw + 0.4, 0.28)
                self._line(tfc, cap, 10, italic=True, color=GREY, align=PP_ALIGN.CENTER, first=True)
            wtxt = BW - fw - 0.40
        else:
            wtxt = BW
        items = list(steps)
        if tip:
            items = items + [("note", tip)]
        items = items + [("gap", 0.04),
                         ("kv", ("Đáp số:", ans))]
        pt, over = self._autofit(items, wtxt, H, base=15.5, low=10.5)
        if over > 0:
            self.warn.append("TRÀN %.2f in — slide %d «%s»" % (over, self.n, head))
        self._emit(s, BX, y + self._voff(items, pt, wtxt, H), wtxt, items, pt, tight=True)
        self._notes(s, note)
        return s

    def table(self, head, headers, rows, sub=None, tag=None, note=None, widths=None,
              foot=None, pt=13.0):
        s = self._slide()
        self._chrome(s, head, sub, tag)
        th = self.th
        y = BY
        H = BH
        fh = 0.0
        if foot:
            fh = wrapn(foot, 13, BW - 0.35) * 13 / 72.0 * 1.3 + 0.22
            H -= fh
        nr, nc = len(rows) + 1, len(headers)
        rowh = min(0.62, max(0.30, H / nr))
        tb = s.shapes.add_table(nr, nc, Inches(BX), Inches(y), Inches(BW),
                                Inches(rowh * nr)).table
        tb.first_row = True
        if widths:
            tot = float(sum(widths))
            for j, wv in enumerate(widths):
                tb.columns[j].width = Inches(BW * wv / tot)
        for r in range(nr):
            tb.rows[r].height = Inches(rowh)
        for j, htxt in enumerate(headers):
            c = tb.cell(0, j)
            c.fill.solid(); c.fill.fore_color.rgb = th["dark"]
            c.margin_left = c.margin_right = Inches(0.07)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            self._run(p, htxt, pt, True, False, WHITE)
        for i, row in enumerate(rows, 1):
            for j, v in enumerate(row):
                c = tb.cell(i, j)
                c.fill.solid()
                c.fill.fore_color.rgb = RGBColor.from_string("FFFFFF") if i % 2 else \
                    RGBColor.from_string(th["tint"])
                c.margin_left = c.margin_right = Inches(0.07)
                c.margin_top = c.margin_bottom = Inches(0.02)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = c.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
                self._run(p, str(v), pt, j == 0, False, INK)
        if foot:
            yy = y + rowh * nr + 0.14
            self._rect(s, BX, yy, 0.045, fh - 0.18, th["primary"])
            tf = self._tb(s, BX + 0.22, yy + 0.01, BW - 0.25, fh)
            self._line(tf, foot, 13, italic=True, color=INK2, first=True)
        self._notes(s, note)
        return s

    def quiz(self, head, q, opts, a, why, sub=None, tag="KIỂM TRA NHANH", note=None, fig=None):
        s = self._slide()
        self._chrome(s, head, sub, tag)
        th = self.th
        qh = wrapn(q, 16, BW - 0.45) * 16 / 72.0 * 1.32
        two_ = all(textw(o, 15) < BW / 2 - 1.0 for o in opts)
        nrow_ = (len(opts) + 1) // 2 if two_ else len(opts)
        hh_ = wrapn(why, 13.5, BW - 1.35) * 13.5 / 72.0 * 1.3
        tot_ = qh + 0.46 + nrow_ * 0.62 + 0.12 + hh_ + 0.30
        y = BY + max(0.0, (BH - tot_) * 0.38)
        self._rect(s, BX, y, BW, qh + 0.32, th["tint"], shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.10)
        tf = self._tb(s, BX + 0.22, y + 0.16, BW - 0.45, qh + 0.05)
        self._line(tf, q, 16, color=INK, first=True)
        y += qh + 0.46
        letters = "ABCD"
        two = two_
        cw = (BW - 0.35) / 2.0 if two else BW
        for i, o in enumerate(opts):
            ok = letters[i] == a
            col = i % 2 if two else 0
            row = i // 2 if two else i
            ox = BX + col * (cw + 0.35)
            oy = y + row * 0.62
            self._rect(s, ox, oy, cw, 0.52,
                       "EAF4EA" if ok else "FFFFFF",
                       line=RGBColor.from_string("2E7D32") if ok else th["soft"],
                       shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.18)
            tf = self._tb(s, ox + 0.18, oy + 0.11, cw - 0.3, 0.34, MSO_ANCHOR.MIDDLE)
            self._line(tf, [("%s. " % letters[i], True, False,
                             RGBColor.from_string("2E7D32") if ok else th["primary"]),
                            (o, False, False, INK)], 14.5, first=True)
        y += ((len(opts) + 1) // 2 if two else len(opts)) * 0.62 + 0.12
        hh = wrapn(why, 13.5, BW - 1.35) * 13.5 / 72.0 * 1.3
        self._rect(s, BX, y, BW, hh + 0.30, "EAF4EA", shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.10)
        tfa = self._tb(s, BX + 0.20, y + 0.15, 1.15, 0.3)
        self._line(tfa, "Đáp án %s" % a, 13.5, bold=True,
                   color=RGBColor.from_string("2E7D32"), first=True)
        tf = self._tb(s, BX + 1.40, y + 0.15, BW - 1.60, hh + 0.05)
        self._line(tf, why, 13.5, color=INK, first=True)
        if y + hh + 0.35 > BY + BH:
            self.warn.append("TRÀN quiz — slide %d «%s»" % (self.n, head))
        self._notes(s, note)
        return s

    def media(self, head, cards, sub=None, tag="TRỰC QUAN HOÁ", note=None, lead=None):
        """cards: [(tên, nguồn, mô tả, 'điều cần quan sát')]"""
        s = self._slide()
        self._chrome(s, head, sub, tag)
        th = self.th
        y = BY
        H = BH
        if lead:
            hh = wrapn(lead, 14.5, BW) * 14.5 / 72.0 * 1.3
            tf = self._tb(s, BX, y, BW, hh + 0.05)
            self._line(tf, lead, 14.5, color=INK2, first=True)
            y += hh + 0.18
            H -= hh + 0.18
        n = len(cards)
        need = 0.0
        for name, src, desc, watch in cards:
            need = max(need,
                       0.34
                       + wrapn(name + "   |   " + src, 14.5, BW - 0.5, True) * 14.5 / 72 * 1.34
                       + wrapn(desc, 13, BW - 0.5) * 13 / 72 * 1.34
                       + wrapn("Cần cho học sinh quan sát: " + watch, 12.5, BW - 0.5) * 12.5 / 72 * 1.34)
        ch = min(need, (H - 0.14 * (n - 1)) / n)
        y += max(0.0, (H - (ch * n + 0.14 * (n - 1))) / 2.0)
        for name, src, desc, watch in cards:
            self._rect(s, BX, y, BW, ch, "FFFFFF", line=th["soft"],
                       shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06)
            self._rect(s, BX, y, 0.075, ch, th["primary"])
            tf = self._tb(s, BX + 0.26, y + 0.12, BW - 0.5, ch - 0.2)
            self._line(tf, [(name, True, False, th["dark"]), ("   |   ", False, False, GREY),
                            (src, False, True, GREY)], 14.5, first=True, space_after=3)
            self._line(tf, desc, 13, color=INK, space_after=3)
            self._line(tf, [("Cần cho học sinh quan sát: ", True, False, th["primary"]),
                            (watch, False, False, INK2)], 12.5)
            y += ch + 0.14
        self._notes(s, note)
        return s

    def wrapup(self, head, items, todo=(), sub=None, note=None):
        s = self._slide()
        self._chrome(s, head, sub, "TỔNG KẾT")
        th = self.th
        lw = BW * 0.58
        rw = BW - lw - 0.45
        pt, _ = self._autofit(items, lw, BH, base=15.5)
        it2 = [("b", t) for t in todo]
        pt2, _ = self._autofit(it2, rw - 0.55, BH - 0.95, base=14)
        cardh = min(BH - 0.15, self._need(it2, pt2, rw - 0.55) + 1.10)
        cy = BY + max(0.0, (BH - cardh) / 2.0)
        self._emit(s, BX, BY + self._voff(items, pt, lw, BH, "center"), lw, items, pt)
        self._rect(s, BX + lw + 0.45, cy, rw, cardh, th["tint"],
                   shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.05)
        tf = self._tb(s, BX + lw + 0.72, cy + 0.24, rw - 0.55, 0.4)
        self._line(tf, "VIỆC CẦN LÀM SAU BUỔI HỌC", 13, bold=True, color=th["dark"], first=True)
        self._emit(s, BX + lw + 0.72, cy + 0.78, rw - 0.55, it2, pt2)
        self._notes(s, note)
        return s

    def save(self, path):
        self.prs.save(path)
        return path


def _norm(it):
    if isinstance(it, str):
        return "b", it
    k = it[0]
    if k in ("b", "p", "sub", "head", "note", "fx", "gap", "rule", "kv", "num"):
        return k, (it[1] if len(it) == 2 else tuple(it[1:]))
    return "b", it[0]


# -------------------------------------------------------------- dựng cả bộ
def build(spec, outdir):
    """spec: dict(file=, title=, chapter=, sub=, meta=, th=, slides=[callable(d)])"""
    d = Deck(spec["title"], spec["chapter"], spec.get("sub", ""), spec.get("meta", ""),
             th=spec.get("th"))
    spec["slides"](d)
    path = os.path.join(outdir, spec["file"])
    d.save(path)
    return path, d.warn, d.n
